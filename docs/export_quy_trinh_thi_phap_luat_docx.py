"""Xuất file Word: Quy trình, quy tắc và hướng dẫn tổ chức thi tìm hiểu pháp luật."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).with_name("QUY_TRINH_QUY_TAC_HUONG_DAN_TO_CHUC_THI_TIM_HIEU_PHAP_LUAT.docx")
RED = RGBColor(0xC0, 0x00, 0x00)
NAVY = RGBColor(0x1F, 0x4E, 0x79)
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x59, 0x59, 0x59)


def set_run_font(run, *, size: int = 13, bold: bool = False, italic: bool = False, color=BLACK, name: str = "Times New Roman"):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)


def shade_cell(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "8FAADC")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_para(doc, text="", *, size=13, bold=False, italic=False, align="left", space_after=6, space_before=0, color=BLACK, first_line=False):
    p = doc.add_paragraph()
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_line:
        pf.first_line_indent = Cm(1.0)
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def add_runs(p, parts: list[tuple[str, dict]]) -> None:
    for text, opts in parts:
        run = p.add_run(text)
        set_run_font(run, **opts)


def heading(doc, text: str, level: int) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(16 if level == 1 else 12)
    pf.space_after = Pt(8)
    pf.keep_with_next = True
    if level == 1:
        set_run_font(p.add_run(text), size=14, bold=True, color=NAVY)
    elif level == 2:
        set_run_font(p.add_run(text), size=13, bold=True, color=NAVY)
    else:
        set_run_font(p.add_run(text), size=13, bold=True, italic=True, color=NAVY)


def bullet(doc, text: str, *, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if bold_prefix:
        set_run_font(p.add_run(bold_prefix), size=13, bold=True)
        set_run_font(p.add_run(text), size=13)
    else:
        set_run_font(p.add_run(text), size=13)


def numbered(doc, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    set_run_font(p.add_run(text), size=13)


def table(doc, headers: list[str], rows: list[list[str]], col_widths: list[float] | None = None) -> None:
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True
    for i, h in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(h), size=12, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade_cell(cell, "1F4E79")
        set_cell_border(cell)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = tbl.rows[r_i + 1].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            set_run_font(p.add_run(val), size=12)
            if r_i % 2 == 1:
                shade_cell(cell, "D6E3F0")
            set_cell_border(cell)
    if col_widths:
        for row in tbl.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    add_para(doc, "", space_after=4)


def quote(doc, text: str) -> None:
    p = add_para(doc, "", align="center", space_before=6, space_after=10)
    set_run_font(p.add_run(text), size=13, bold=True, italic=True, color=NAVY)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(hp.add_run("CÔNG ĐOÀN TRƯỜNG ĐẠI HỌC SƯ PHẠM KỸ THUẬT HƯNG YÊN"), size=10, italic=True, color=GRAY)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(fp.add_run("Quy trình — Quy tắc — Hướng dẫn tổ chức thi và thi Tìm hiểu pháp luật  |  Trang "), size=9, italic=True, color=GRAY)
    # PAGE field
    run = fp.add_run()
    set_run_font(run, size=9, italic=True, color=GRAY)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)

    # ===== Bìa =====
    add_para(doc, "CÔNG ĐOÀN VIỆT NAM", size=13, bold=True, align="center", space_after=0, color=RED)
    add_para(doc, "CÔNG ĐOÀN TRƯỜNG ĐẠI HỌC SƯ PHẠM KỸ THUẬT HƯNG YÊN", size=13, bold=True, align="center", space_after=2, color=RED)
    line = add_para(doc, "", align="center", space_after=14)
    set_run_font(line.add_run("————————"), size=12, color=RED)

    add_para(doc, "QUY TRÌNH — QUY TẮC", size=18, bold=True, align="center", space_before=24, space_after=4, color=NAVY)
    add_para(doc, "VÀ HƯỚNG DẪN TỔ CHỨC THI, THI", size=18, bold=True, align="center", space_after=4, color=NAVY)
    add_para(doc, "TÌM HIỂU PHÁP LUẬT", size=20, bold=True, align="center", space_after=16, color=NAVY)

    add_para(
        doc,
        "Áp dụng trên Website Công đoàn UTEHY — Tiện ích số / Phổ biến pháp luật",
        size=12,
        italic=True,
        align="center",
        space_after=6,
        color=GRAY,
    )
    add_para(
        doc,
        f"Ban hành kèm theo việc triển khai chức năng thi trắc nghiệm trên website  ·  {date.today().strftime('%d/%m/%Y')}",
        size=12,
        italic=True,
        align="center",
        space_after=18,
        color=GRAY,
    )

    add_para(doc, "NỘI DUNG GỒM BA PHẦN", size=12, bold=True, align="center", space_after=4, color=NAVY)
    add_para(doc, "Phần I.  Quy trình tổ chức thi tìm hiểu pháp luật", size=13, align="center", space_after=2)
    add_para(doc, "Phần II.  Quy tắc cuộc thi (đề thi, thời gian, tính điểm, xếp hạng)", size=13, align="center", space_after=2)
    add_para(doc, "Phần III.  Hướng dẫn tổ chức thi (cán bộ) và hướng dẫn thi (đoàn viên)", size=13, align="center", space_after=18)

    # ===== PHẦN I =====
    heading(doc, "PHẦN I. QUY TRÌNH TỔ CHỨC THI TÌM HIỂU PHÁP LUẬT", 1)

    heading(doc, "1. Mục đích", 2)
    bullet(doc, "Phổ biến, giáo dục pháp luật cho công đoàn viên theo từng đợt (quý, chuyên đề hoặc đột xuất).")
    bullet(doc, "Tổ chức thi trắc nghiệm công bằng: mỗi người nhận một đề được hệ thống trộn từ ngân hàng câu hỏi; không xem đáp án khi đang làm bài chính thức.")
    bullet(doc, "Tổng hợp kết quả để khen thưởng cá nhân và đánh giá mức độ tham gia của từng công đoàn bộ phận.")

    heading(doc, "2. Phạm vi áp dụng", 2)
    bullet(doc, "công đoàn viên đã có tài khoản đăng nhập và đã được gắn với hồ sơ trong danh bạ (để hệ thống xác định công đoàn bộ phận).", bold_prefix="Đối tượng dự thi: ")
    bullet(doc, "công khai, không cần đăng nhập.", bold_prefix="Đọc tài liệu: ")
    bullet(doc, "bắt buộc đăng nhập.", bold_prefix="Thi thử và thi chính thức: ")
    bullet(doc, "Không áp dụng cho chức năng Khảo sát (ý kiến ẩn danh, không chấm điểm).")

    heading(doc, "3. Trách nhiệm các bên", 2)
    table(
        doc,
        ["Vai trò", "Trách nhiệm"],
        [
            [
                "Ban chấp hành Công đoàn trường",
                "Quyết định chủ đề, lịch thi thử / thi chính thức, mức đạt, số câu, thời lượng, thể lệ tính điểm và công bố kết quả.",
            ],
            [
                "Cán bộ quản trị website (ADMIN / UNION_CLERK)",
                "Tạo đợt, nhập tài liệu, nhập ngân hàng câu hỏi, cấu hình cửa sổ thi, mở–khóa thi, xuất thống kê.",
            ],
            [
                "Chủ tịch / ủy viên công đoàn bộ phận",
                "Đôn đốc đoàn viên đọc tài liệu, thi thử, thi chính thức; đối chiếu danh sách đoàn viên đã gắn tài khoản.",
            ],
            [
                "Công đoàn viên",
                "Đọc tài liệu; thi thử (không tính xếp hạng); thi chính thức trong thời gian quy định, đúng tài khoản đã gắn hồ sơ.",
            ],
        ],
        [5.5, 11.0],
    )

    heading(doc, "4. Lịch trình một đợt thi", 2)
    add_para(
        doc,
        "Ban chấp hành chốt lịch trước khi mở công khai. Nên khóa thi thử trước khi mở thi chính thức. Nếu hai cửa sổ chồng nhau, đoàn viên vẫn chọn đúng nút Thi thử hoặc Thi chính thức.",
        align="justify",
        first_line=True,
    )
    table(
        doc,
        ["Giai đoạn", "Việc cần làm", "Hệ thống"],
        [
            [
                "A. Chuẩn bị",
                "Soạn tài liệu, nhập ngân hàng, rà đáp án, cấu hình số câu / phút / điểm đạt.",
                "Đợt ở trạng thái Nháp, hoặc đã công khai tài liệu nhưng Chưa mở thi.",
            ],
            [
                "B. Công bố tài liệu",
                "Thông báo trên website, Zalo / email nội bộ.",
                "Bật Công khai — hiện trên trang tiện ích số.",
            ],
            [
                "C. Thi thử",
                "Đoàn viên ôn tập, làm quen giao diện, xem đáp án sau khi nộp.",
                "Điền Mở / Khóa thi thử, Số lần thi thử (gợi ý 3–5 lần).",
            ],
            [
                "D. Thi chính thức",
                "Khóa thi thử (hoặc để hết hạn). Thông báo giờ mở–khóa rõ ràng.",
                "Bật Đang mở thi; điền giờ mở–khóa chính thức. Chỉ lượt chính thức mới tính điểm.",
            ],
            [
                "E. Khóa sổ",
                "Đến giờ khóa hoặc khóa sớm theo quyết định Ban chấp hành.",
                "Tắt Đang mở thi hoặc hết hạn khóa. Không bắt đầu lượt chính thức mới.",
            ],
            [
                "F. Tổng hợp — công bố",
                "Xem Kết quả, xuất CSV, trình Ban chấp hành công bố.",
                "Bảng xếp hạng cá nhân, điểm công đoàn bộ phận, CSV cá nhân / đơn vị.",
            ],
        ],
        [3.5, 6.5, 6.5],
    )

    heading(doc, "5. Checklist trước giờ mở thi chính thức", 2)
    for item in [
        "Tài liệu đã xuất bản; đường dẫn kiểm tra được trên máy khác.",
        "Ngân hàng đủ câu; số câu mỗi đề không vượt quá số câu ngân hàng.",
        "Xáo câu hỏi = Có; Xáo lựa chọn = Có.",
        "Thời lượng, điểm đạt, số lần thi chính thức đã chốt.",
        "Cửa sổ thi thử đã khóa hoặc đã hết hạn (nếu Ban chấp hành yêu cầu).",
        "Cửa sổ thi chính thức đúng lịch; trạng thái Đang mở thi.",
        "Đoàn viên đã gắn tài khoản và đúng công đoàn bộ phận.",
        "Đã chạy thử 1 lượt thi thử và kiểm tra đồng hồ đếm ngược, tự nộp hết giờ.",
    ]:
        bullet(doc, item)

    heading(doc, "6. Lưu hồ sơ sau khóa sổ", 2)
    bullet(doc, "CSV cá nhân (mọi lượt; cột Loại = Thi thử / Chính thức, bộ phận, %).")
    bullet(doc, "CSV đơn vị (hạng, tham gia, điểm trung bình, tỷ lệ đạt, điểm đơn vị).")
    bullet(doc, "Bảng xếp hạng (in hoặc lưu ảnh) nếu cần công bố.")
    add_para(
        doc,
        "Không công bố đáp án ngân hàng trên kênh rộng nếu còn dùng lại cho đợt sau.",
        align="justify",
        first_line=True,
    )

    # ===== PHẦN II =====
    heading(doc, "PHẦN II. QUY TẮC CUỘC THI", 1)
    add_para(
        doc,
        "Các quy tắc dưới đây do hệ thống thực thi. Cán bộ không được tắt, trừ khi Ban chấp hành cho phép bằng văn bản.",
        align="justify",
        first_line=True,
    )

    heading(doc, "1. Quy tắc ra đề", 2)
    numbered(doc, "Đề thi được trộn từ ngân hàng câu hỏi của đợt. Mọi câu nhập tại màn hình Câu hỏi là ngân hàng. Số câu trên mỗi đề do Ban chấp hành quyết định và có thể ít hơn số câu trong ngân hàng. Khi đoàn viên bấm thi, máy chủ lấy ngẫu nhiên đúng số câu đã cấu hình. Để trống số câu mỗi đề nghĩa là dùng hết ngân hàng.")
    numbered(doc, "Trong mỗi bài thi, hệ thống trộn thứ tự câu hỏi (Xáo câu hỏi = Có). Thứ tự đã khóa khi bắt đầu lượt; F5 / tải lại trang không đổi đề.")
    numbered(doc, "Trong mỗi câu, hệ thống trộn thứ tự đáp án (Xáo lựa chọn = Có). Máy chủ chấm theo chỉ số đáp án gốc, không theo thứ tự hiển thị trên màn hình.")
    numbered(doc, "Thời gian làm bài tùy biến theo từng đợt (Thời gian mỗi lượt, đơn vị phút). Số câu trên đề tùy biến (Số câu mỗi đề).")

    heading(doc, "2. Quy tắc thời gian mở — khóa — thi thử", 2)
    bullet(doc, "phải bật Đang mở thi và nằm trong khoảng Mở thi chính thức từ → Khóa thi chính thức. Ngoài khoảng này không bắt đầu được lượt chính thức mới.", bold_prefix="Thi chính thức: ")
    bullet(doc, "nằm trong khoảng Mở thi thử từ → Khóa thi thử; có số lần thi thử riêng. Lượt thi thử không tính vào xếp hạng cá nhân và điểm đơn vị. Sau khi nộp thi thử, hệ thống hiện đáp án để ôn.", bold_prefix="Thi thử: ")
    bullet(doc, "Đến giờ khóa, hệ thống từ chối lượt mới. Lượt đang làm vẫn được nộp trong thời lượng còn lại, cộng 60 giây dự phòng phía máy chủ (tránh lệch đồng hồ / mạng chậm). Có thể khóa sớm bằng cách chuyển trạng thái về Chưa mở thi.")

    heading(doc, "3. Đồng hồ đếm ngược khi đoàn viên thi", 2)
    add_para(
        doc,
        "Hệ thống đã có đồng hồ đếm ngược trên trang thi (cả thi thử và thi chính thức).",
        align="justify",
        first_line=True,
    )
    bullet(doc, "Góc trên trang thi hiện “Còn lại mm:ss”, tính từ lúc bắt đầu lượt theo thời gian mỗi lượt đã cấu hình.")
    bullet(doc, "Khi còn từ 60 giây trở xuống, đồng hồ chuyển màu cảnh báo.")
    bullet(doc, "Hết giờ, hệ thống tự động nộp phần đã làm. Lượt hết giờ vẫn được chấm trên các câu đã trả lời.")
    bullet(doc, "F5 không reset giờ: lượt đang làm được tiếp tục với thời gian còn lại (mốc hết giờ đã khóa lúc bắt đầu).")
    bullet(doc, "Câu trả lời được lưu tạm trong lúc làm bài; mất mạng ngắn không mất hết bài nếu đã kịp ghi nhận.")

    heading(doc, "4. Quy tắc tính điểm cá nhân", 2)
    bullet(doc, "Mỗi câu đúng = 1 điểm; câu sai hoặc bỏ trống = 0 điểm.")
    bullet(doc, "Tỷ lệ % = (số câu đúng / số câu trên đề) × 100.")
    bullet(doc, "Đạt khi % ≥ ngưỡng Điểm đạt (%) của đợt (mặc định 70%).")
    bullet(doc, "Chỉ lượt thi chính thức mới tính thành tích. Thi thử không tính.")
    bullet(doc, "Nếu được phép nhiều lần thi chính thức: lấy lượt tốt nhất của người đó (% cao hơn; nếu cùng % thì lấy lượt nộp sớm hơn).")

    heading(doc, "5. Quy tắc xếp hạng cá nhân khi cùng điểm", 2)
    add_para(
        doc,
        "Bảng Xếp hạng cá nhân chỉ lấy lượt thi chính thức tốt nhất của mỗi người.",
        align="justify",
        first_line=True,
    )
    numbered(doc, "Người có tỷ lệ % cao hơn xếp trên.")
    numbered(doc, "Khi cùng điểm số / cùng %, người nộp bài sớm hơn xếp trên (căn cứ giờ nộp trên máy chủ).")
    numbered(doc, "Số hạng là 1, 2, 3… liên tiếp. Hai người cùng điểm không chia chung một hạng.")
    add_para(
        doc,
        "Ví dụ: A và B cùng 18/20 câu (90%). A nộp lúc 09:05, B nộp lúc 09:08 → A hạng 1, B hạng 2.",
        italic=True,
        align="justify",
        space_before=4,
    )

    heading(doc, "6. Quy tắc tính điểm và xếp hạng công đoàn bộ phận", 2)
    add_para(
        doc,
        "Chỉ tính công đoàn viên đã gắn tài khoản với hồ sơ thuộc bộ phận đó. Mẫu số “có tài khoản” lấy từ danh bạ Công đoàn bộ phận.",
        align="justify",
        first_line=True,
    )
    bullet(doc, "số người đã nộp bài chính thức / số đoàn viên có tài khoản trong bộ phận × 100.", bold_prefix="Tỷ lệ tham gia = ")
    bullet(doc, "trung bình % của các lượt chính thức tốt nhất trong bộ phận (0 nếu chưa ai nộp).", bold_prefix="Điểm trung bình = ")
    bullet(doc, "số người đạt / số người đã nộp chính thức × 100 (0 nếu chưa ai nộp).", bold_prefix="Tỷ lệ đạt = ")
    quote(doc, "Điểm đơn vị = 50% × Điểm trung bình  +  30% × Tỷ lệ tham gia  +  20% × Tỷ lệ đạt")
    add_para(
        doc,
        "Xếp hạng đơn vị theo điểm đơn vị giảm dần; nếu hòa điểm đơn vị thì ưu tiên tỷ lệ tham gia cao hơn. Người thi nhưng chưa gắn công đoàn bộ phận vẫn có trên bảng cá nhân (ghi “Chưa gắn bộ phận”) và không cộng vào điểm đơn vị. Thi thử không tham gia bất kỳ chỉ số đơn vị nào.",
        align="justify",
        first_line=True,
    )

    heading(doc, "7. Các hành vi không hợp lệ", 2)
    bullet(doc, "Thi hộ, chia sẻ tài khoản, cố tình mở nhiều trình duyệt để lấy nhiều đề rồi chọn đề dễ.")
    bullet(doc, "Công bố đáp án ngân hàng trong thời gian thi chính thức còn mở.")
    bullet(doc, "Yêu cầu cán bộ xóa lượt thi để thi lại, trừ khi Ban chấp hành quyết định bằng văn bản (thường phải mở đợt mới).")
    add_para(
        doc,
        "Mỗi lượt thi lưu riêng thứ tự câu và thứ tự đáp án. Hai người không nhận cùng một đề / cùng thứ tự đáp án nếu đã bật xáo.",
        align="justify",
        first_line=True,
    )

    # ===== PHẦN III =====
    heading(doc, "PHẦN III. HƯỚNG DẪN TỔ CHỨC THI VÀ HƯỚNG DẪN THI", 1)

    heading(doc, "A. Hướng dẫn tổ chức thi (dành cho cán bộ)", 2)
    add_para(
        doc,
        "Đường dẫn quản trị: menu Phổ biến pháp luật ( /legal-education-campaigns ). Đường dẫn công khai: Tiện ích số / Phổ biến pháp luật.",
        align="justify",
        first_line=True,
    )

    heading(doc, "Bước 1. Tạo đợt phổ biến", 3)
    numbered(doc, "Vào Phổ biến pháp luật → Thêm đợt.")
    numbered(doc, "Nhập tiêu đề, nhãn thời gian (ví dụ Quý III năm 2026), tóm tắt.")
    numbered(doc, "Cấu hình bài thi: Thời gian mỗi lượt (phút); Số câu mỗi đề (ví dụ 20 câu từ ngân hàng 45 câu; để trống = dùng hết ngân hàng); Điểm đạt (%); Số lần thi chính thức (thường 1); Hiện đáp án sau khi nộp = Không đối với thi chính thức (thi thử vẫn hiện đáp án); Xáo câu hỏi = Có; Xáo lựa chọn = Có; Trạng thái thi = Chưa mở thi cho đến khi vào giai đoạn D; điền Mở / Khóa thi chính thức và Mở / Khóa thi thử, Số lần thi thử.")
    numbered(doc, "Xuất bản tài liệu: có thể công khai tài liệu trước, chưa mở thi. Lưu. Hệ thống chuyển sang màn hình tài liệu.")

    heading(doc, "Bước 2. Đăng tài liệu ôn tập", 3)
    add_para(
        doc,
        "Tại Tài liệu của đợt: thêm tóm tắt, file PDF (nếu có), bật xuất bản từng bài. Đoàn viên đọc không cần đăng nhập.",
        align="justify",
        first_line=True,
    )

    heading(doc, "Bước 3. Xây dựng ngân hàng câu hỏi", 3)
    bullet(doc, "Mỗi câu: nội dung + ít nhất 2 lựa chọn (mỗi lựa chọn một dòng) + chọn đáp án đúng.")
    bullet(doc, "Nên nhập nhiều hơn số câu trên mỗi đề để việc lấy ngẫu nhiên có ý nghĩa (ví dụ ngân hàng 45, mỗi đề 20).")
    bullet(doc, "Không công khai đáp án trên trang đoàn viên khi đang thi chính thức. Chốt ngân hàng trước giờ mở chính thức; tránh xóa / sửa câu đã có người thi.")

    heading(doc, "Bước 4. Rà soát trước khi mở", 3)
    bullet(doc, "Làm một lượt thi thử bằng tài khoản thật trên trang công khai.")
    bullet(doc, "Kiểm tra: số câu đúng bằng cấu hình; câu và đáp án đã xáo; đồng hồ đếm ngược; nộp hết giờ; kết quả thi thử không lên bảng xếp hạng chính thức.")
    bullet(doc, "Xác nhận đoàn viên đã được gắn tài khoản với đúng công đoàn bộ phận (Danh bạ công đoàn viên). Nếu thiếu liên kết, người đó vẫn thi được nhưng đơn vị không được tính.")

    heading(doc, "Bước 5. Mở thi thử", 3)
    numbered(doc, "Sửa đợt: điền giờ mở–khóa thi thử, số lần thi thử.")
    numbered(doc, "Công bố đường dẫn đợt. Đoàn viên bấm Thi thử (cần đăng nhập).")
    numbered(doc, "Theo dõi tab Tất cả lượt thi — cột Thi thử.")

    heading(doc, "Bước 6. Mở và khóa thi chính thức", 3)
    numbered(doc, "Tắt hoặc để hết hạn thi thử (khuyến nghị).")
    numbered(doc, "Bật Đang mở thi, điền Mở thi chính thức từ và Khóa thi chính thức.")
    numbered(doc, "Thông báo giờ khóa rõ ràng. Đến giờ khóa, hệ thống từ chối lượt chính thức mới.")
    numbered(doc, "Có thể khóa sớm bằng cách chuyển Trạng thái thi về Chưa mở thi.")

    heading(doc, "Bước 7. Tổng hợp — xuất số liệu", 3)
    add_para(doc, "Vào Kết quả của đợt:", align="justify")
    bullet(doc, "hạng, họ tên, công đoàn bộ phận, điểm, %, đạt / không đạt, giờ nộp. Chỉ lượt chính thức tốt nhất.", bold_prefix="Xếp hạng cá nhân: ")
    bullet(doc, "số có tài khoản, số đã nộp, số đạt, tỷ lệ tham gia, điểm trung bình, tỷ lệ đạt, điểm đơn vị.", bold_prefix="Công đoàn bộ phận: ")
    bullet(doc, "gồm thi thử và chính thức, phục vụ đối soát.", bold_prefix="Tất cả lượt thi: ")
    bullet(doc, "CSV cá nhân / CSV đơn vị để lưu hồ sơ, báo cáo Ban chấp hành.", bold_prefix="Xuất file: ")

    heading(doc, "Xử lý sự cố thường gặp khi tổ chức", 3)
    table(
        doc,
        ["Tình huống", "Cách xử lý"],
        [
            [
                "“Bài thi này chưa mở hoặc đã kết thúc”",
                "Kiểm tra Đang mở thi và giờ mở–khóa; đồng hồ máy chủ.",
            ],
            [
                "“Chưa đến hoặc đã hết thời gian thi thử”",
                "Ngoài cửa sổ thi thử; dùng đúng nút Thi thử.",
            ],
            [
                "“Đã hết số lần thi”",
                "Hết số lần thi chính thức hoặc số lần thi thử. Không xóa lượt để cho thi lại trừ khi Ban chấp hành quyết định tạo đợt mới.",
            ],
            [
                "Người thi không hiện ở bảng đơn vị",
                "Hồ sơ chưa gắn tài khoản hoặc chưa gán công đoàn bộ phận. Gắn xong, mở lại Kết quả — không cần thi lại.",
            ],
            [
                "Hai người cùng đề / cùng thứ tự đáp án",
                "Không xảy ra nếu đã bật xáo; mỗi lượt lưu thứ tự câu và đáp án riêng.",
            ],
            [
                "Sửa / xóa câu hỏi khi đã có người thi",
                "Tránh xóa câu đã dùng. Sửa nội dung có thể làm lệch bài đã nộp. Chốt ngân hàng trước giờ mở chính thức.",
            ],
        ],
        [6.5, 10.0],
    )

    heading(doc, "B. Hướng dẫn thi (dành cho công đoàn viên)", 2)
    heading(doc, "1. Chuẩn bị", 3)
    numbered(doc, "Có tài khoản đăng nhập website Công đoàn và đã được gắn với hồ sơ công đoàn viên (đúng công đoàn bộ phận).")
    numbered(doc, "Máy tính hoặc điện thoại có mạng ổn định; dùng trình duyệt cập nhật; không mở nhiều thẻ thi cùng lúc.")
    numbered(doc, "Đọc hết tài liệu tóm tắt / PDF của đợt trước khi bấm thi.")

    heading(doc, "2. Các bước làm bài", 3)
    numbered(doc, "Vào Tiện ích số → Phổ biến pháp luật → chọn đợt đang mở.")
    numbered(doc, "Trong thời gian thi thử: bấm Thi thử. Được làm nhiều lần theo cấu hình; sau khi nộp được xem đáp án; không tính thành tích, không tính điểm đơn vị.")
    numbered(doc, "Trong thời gian thi chính thức: bấm Thi chính thức. Đề được lấy ngẫu nhiên từ ngân hàng, đã trộn thứ tự câu và đáp án.")
    numbered(doc, "Theo dõi đồng hồ “Còn lại mm:ss” ở góc trên. Dưới 1 phút sẽ cảnh báo. Hết giờ hệ thống tự nộp, không cần thao tác thêm.")
    numbered(doc, "Chọn đáp án từng câu; có thể quay lại câu trước. Thanh số câu cho biết câu nào đã trả lời.")
    numbered(doc, "Nộp bài trước khi hết giờ nếu đã xong. Sau khi nộp thi chính thức, màn hình hiện Đạt / Không đạt và số câu đúng. Đáp án chi tiết chỉ hiện nếu Ban chấp hành bật (thường tắt với thi chính thức).")

    heading(doc, "3. Lưu ý khi thi", 3)
    bullet(doc, "Đăng nhập đúng tài khoản đã gắn hồ sơ để kết quả vào đúng công đoàn bộ phận.")
    bullet(doc, "Không thoát giữa chừng nếu chưa muốn kết thúc; đóng nhầm trang rồi vào lại vẫn tiếp tục đúng lượt và đúng thời gian còn lại.")
    bullet(doc, "Không nhờ người khác thi hộ, không chia sẻ tài khoản.")
    bullet(doc, "Nếu báo “chưa mở / đã khóa / hết số lần thi”: chờ đúng cửa sổ thời gian hoặc liên hệ công đoàn bộ phận, không tự tạo tài khoản thứ hai.")

    heading(doc, "4. Xem kết quả của mình", 3)
    add_para(
        doc,
        "Ngay sau khi nộp, đoàn viên thấy điểm lượt vừa làm. Bảng xếp hạng cá nhân và điểm công đoàn bộ phận do Ban chấp hành công bố sau khi khóa sổ, căn cứ số liệu trên trang Kết quả của cán bộ quản trị.",
        align="justify",
        first_line=True,
    )

    # Kết =====
    heading(doc, "ĐIỀU KHOẢN THI HÀNH", 1)
    add_para(
        doc,
        "Quy trình, quy tắc và hướng dẫn này được thực hiện thống nhất trên Website Công đoàn Trường Đại học Sư phạm Kỹ thuật Hưng Yên. Khi Ban chấp hành điều chỉnh thể lệ một đợt cụ thể (số câu, thời lượng, mức đạt, lịch thi thử / thi chính thức), cán bộ cấu hình đúng trên đợt đó; các nguyên tắc trộn đề, trộn đáp án, đồng hồ đếm ngược, xếp hạng khi cùng điểm và công thức điểm đơn vị giữ nguyên trừ khi có văn bản sửa đổi.",
        align="justify",
        first_line=True,
    )

    add_para(doc, "", space_before=24)
    sign = doc.add_table(rows=1, cols=2)
    sign.autofit = True
    left, right = sign.rows[0].cells
    left.text = ""
    right.text = ""
    p1 = right.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p1.add_run("BAN CHẤP HÀNH CÔNG ĐOÀN TRƯỜNG"), size=13, bold=True)
    p2 = right.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p2.add_run("(Ký, ghi rõ họ tên)"), size=12, italic=True, color=GRAY)

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
